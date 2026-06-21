import argparse
import json
import os
import sys
from pathlib import Path
from typing import Any


def write_json(path: Path, payload: dict) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")


def read_xlsx_listings(path: Path, job_code: str | None) -> list[dict[str, Any]]:
    if not path.exists():
        return []

    from openpyxl import load_workbook

    workbook = load_workbook(path)
    sheet = workbook.active
    rows = list(sheet.iter_rows(values_only=True))
    if len(rows) < 2:
        return []

    headers = [str(value or "") for value in rows[0]]
    listings = []
    for row in rows[1:]:
        record = dict(zip(headers, row))
        title = record.get("Название")
        price = record.get("Цена")
        url = record.get("URL")
        if not title or not isinstance(price, int | float) or not url:
            continue
        listings.append(
            {
                "source": "avito_duff89",
                "job_code": job_code,
                "external_id": str(url).rstrip("/").split("/")[-1],
                "title": str(title),
                "description": str(record.get("Описание") or ""),
                "price": int(price),
                "currency": "RUB",
                "url": str(url),
                "location": record.get("Адрес"),
                "published_at": None,
                "raw_payload": record,
            }
        )
    return listings


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("--repo", default="/tmp/duff89-parser-avito")
    parser.add_argument("--url", default=os.environ.get("AVITO_DIAGNOSTIC_SEARCH_URL"))
    parser.add_argument("--job-code", default=os.environ.get("AVITO_DIAGNOSTIC_JOB_CODE"))
    parser.add_argument("--job-dir", default=os.environ.get("AVITO_DIAGNOSTIC_JOB_DIR", "."))
    parser.add_argument("--timeout", type=int, default=20)
    args = parser.parse_args()

    if not args.url:
        raise SystemExit("--url is required")

    repo = Path(args.repo)
    job_dir = Path(args.job_dir)
    output_path = job_dir / "duff89_probe_report.json"
    if not repo.exists():
        payload = {
            "status": "missing",
            "repo": str(repo),
            "error": "Duff89/parser_avito repo not found",
        }
        write_json(output_path, payload)
        print(json.dumps(payload, ensure_ascii=False, indent=2))
        return 2

    workspace = job_dir / "duff89_probe_workspace"
    workspace.mkdir(parents=True, exist_ok=True)
    (workspace / "logs").mkdir(exist_ok=True)
    (workspace / "result").mkdir(exist_ok=True)
    result_path = workspace / "result" / "avito.xlsx"
    if result_path.exists():
        result_path.unlink()

    sys.path.insert(0, str(repo))
    os.chdir(workspace)

    try:
        from dto import AvitoConfig
        from parser_cls import AvitoParse

        config = AvitoConfig(
            urls=[args.url],
            count=1,
            proxy_string="",
            proxy_change_url="",
            keys_word_white_list=[],
            keys_word_black_list=[],
            seller_black_list=[],
            max_price=999999999,
            min_price=0,
            geo="",
            pause_general=1,
            pause_between_links=1,
            max_age=0,
            max_count_of_retry=1,
            ignore_reserv=True,
            ignore_promotion=False,
            one_time_start=True,
            one_file_for_link=False,
            parse_views=False,
            save_xlsx=True,
            use_webdriver=False,
            use_bypass_api=False,
            cookies_api_key="",
            output_dir=Path("result"),
            use_own_cookies=False,
            parse_phone=False,
            proxy_notifier="",
            tg_only_text=False,
            retry_delay=1,
            timeout=args.timeout,
            block_threshold=1,
            tg_token="",
            tg_chat_id=[],
            vk_token="",
            vk_user_id=[],
        )
        parser_instance = AvitoParse(config)
        parser_instance.parse()
        listings = read_xlsx_listings(result_path, args.job_code)
        write_json(job_dir / "duff89_normalized_listings.json", {"items": listings})
        payload = {
            "status": "success"
            if parser_instance.good_request_count > 0 and parser_instance.bad_request_count == 0 and listings
            else "failed",
            "provider": "Duff89/parser_avito",
            "repo": str(repo),
            "job_code": args.job_code,
            "url": args.url,
            "good_request_count": parser_instance.good_request_count,
            "bad_request_count": parser_instance.bad_request_count,
            "listings_count": len(listings),
            "result_path": str(result_path),
            "workspace": str(workspace),
        }
    except Exception as exc:
        payload = {
            "status": "failed",
            "provider": "Duff89/parser_avito",
            "repo": str(repo),
            "job_code": args.job_code,
            "url": args.url,
            "error": f"{type(exc).__name__}: {exc}",
            "workspace": str(workspace),
        }

    write_json(output_path, payload)
    print(json.dumps(payload, ensure_ascii=False, indent=2))
    return 0 if payload["status"] == "success" else 1


if __name__ == "__main__":
    raise SystemExit(main())
